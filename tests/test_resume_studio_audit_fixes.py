"""
Regression tests for Resume Studio audit fixes:
1. Cover letter persistence in update_resume SQL query
2. Draft resume commit without opportunity foreign key crash (opportunity_id=None)
3. Download / export query parameter guest token authentication
4. PDF and DOCX file parser resilience (pypdf, python-docx, pdfplumber)
5. Self-anchored master resume validation
"""
import io
import pytest
import httpx
from fastapi.testclient import TestClient
from backend.main import app
from backend.models_career import (
    CommitExtractedResumeRequest,
    ResumeContent,
    ResumeOut,
    ResumeValidationReport,
    ResumeValidationStatus,
    SaveResumeRequest,
)
from backend.services.career_file_parser import (
    extract_resume_from_file_bytes,
    is_docx_parser_available,
    is_pdf_parser_available,
)
from backend.services.career_resume import validate_resume_against_master

client = TestClient(app)

def test_pdf_and_docx_parsers_available():
    """Verify pypdf/pdfplumber and python-docx are installed and functional."""
    assert is_pdf_parser_available() is True, "PDF parser must be available"
    assert is_docx_parser_available() is True, "DOCX parser must be available"


def test_docx_file_extraction_fidelity():
    """Verify DOCX resume text extraction and structured parsing."""
    import docx
    doc = docx.Document()
    doc.add_heading("Alice Walker", 0)
    doc.add_paragraph("alice@example.com | Seattle, WA | +1 555-0144")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("University of Washington - BS Computer Science, 2023")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, React, TypeScript, Docker, PostgreSQL")
    
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    result = extract_resume_from_file_bytes(docx_bytes, filename="alice_resume.docx")
    assert result.success is True
    assert result.extracted_content is not None
    assert result.extracted_content.full_name == "Alice Walker"
    assert "Python" in result.extracted_content.skills
    assert "PostgreSQL" in result.extracted_content.skills


def test_commit_extracted_draft_no_fk_violation():
    """Verify committing separate draft does not crash on foreign key constraint."""
    from backend.dependencies import generate_guest_token
    gid, token = generate_guest_token()
    headers = {"x-guest-token": token}

    payload = {
        "content": {
            "full_name": "Test Candidate",
            "email": "candidate@example.com",
            "skills": ["Python", "FastAPI"],
        },
        "target": "draft",
    }
    resp = client.post("/api/career/resumes/commit-extracted", json=payload, headers=headers)
    assert resp.status_code == 200, f"Expected 200, got {resp.status_code}: {resp.text}"
    data = resp.json()
    assert data["is_master"] is False
    assert data["opportunity_id"] is None


def test_cover_letter_persisted_in_update_resume():
    """Verify cover letter is preserved and returned on update_resume."""
    from backend.dependencies import generate_guest_token
    gid, token = generate_guest_token()
    headers = {"x-guest-token": token}

    # 1. Create a draft resume
    draft_payload = {
        "content": {
            "full_name": "Cover Letter Tester",
            "email": "tester@example.com",
            "skills": ["Python", "AsyncIO"],
        },
        "target": "draft",
    }
    resp_draft = client.post("/api/career/resumes/commit-extracted", json=draft_payload, headers=headers)
    assert resp_draft.status_code == 200
    draft_id = resp_draft.json()["id"]

    # 2. Update with cover letter
    letter_text = "Dear Team,\nI am writing to express my strong interest in the role.\nBest regards,\nTester"
    update_payload = {
        "title": "Updated Draft with Letter",
        "content": {
            "full_name": "Cover Letter Tester",
            "email": "tester@example.com",
            "skills": ["Python", "AsyncIO"],
        },
        "cover_letter": letter_text,
    }
    resp_update = client.put(f"/api/career/resumes/{draft_id}", json=update_payload, headers=headers)
    assert resp_update.status_code == 200
    assert resp_update.json()["cover_letter"] == letter_text

    # 3. Fetch again and confirm persistence
    resp_get = client.get(f"/api/career/resumes/{draft_id}", headers=headers)
    assert resp_get.status_code == 200
    assert resp_get.json()["cover_letter"] == letter_text


def test_export_endpoints_authenticated_via_query_param():
    """Verify HTML and DOCX export endpoints accept guest_token in query parameter without headers."""
    from backend.dependencies import generate_guest_token
    gid, token = generate_guest_token()
    headers = {"x-guest-token": token}

    # Fetch master resume to get valid id
    resp_master = client.get("/api/career/resumes/master", headers=headers)
    assert resp_master.status_code == 200
    res_id = resp_master.json()["id"]

    # Request HTML export with query param token and NO headers
    resp_html = client.get(f"/api/career/resumes/{res_id}/export/html?guest_token={token}&template=minimal")
    assert resp_html.status_code == 200
    assert "resume-template-minimal" in resp_html.text

    # Request DOCX export with query param token and NO headers
    resp_docx = client.get(f"/api/career/resumes/{res_id}/export/docx?guest_token={token}")
    assert resp_docx.status_code == 200
    assert len(resp_docx.content) > 1000
    assert resp_docx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_master_validation_self_anchors():
    """Verify validating master resume evaluates against its own content without false unsupported claims."""
    content = ResumeContent(
        full_name="Verified Master Candidate",
        email="verified@example.com",
        skills=["Python", "PostgreSQL", "Docker"],
        education=[],
        projects=[],
    )
    report = validate_resume_against_master(content, content)
    assert report.is_valid is False  # Missing education flag
    # But all skills are verified against master (0 unsupported)
    assert len(report.unsupported_claims) == 0
    assert any("All 3 listed skills verified" in p for p in report.passed_checks)


def test_rhythm_resume_extraction_fidelity():
    """Verify extraction accuracy on the user's attached resume (projects, honors, skills)."""
    import os
    docx_path = r"C:\Users\User\Downloads\RHYTHM_BACHANI_FlowCV_Resume_2026-10-17.docx"
    if not os.path.exists(docx_path):
        pytest.skip("Local sample resume file not found in test environment")

    with open(docx_path, "rb") as f:
        res = extract_resume_from_file_bytes(f.read(), "RHYTHM_BACHANI_FlowCV_Resume_2026-10-17.docx")

    assert res.success is True
    c = res.extracted_content
    assert c is not None

    # 1. Projects validation: Exactly 2 projects with proper titles, not bullet fragments
    assert len(c.projects) == 2
    p1 = c.projects[0]
    assert "CareMaa" in p1.title
    assert "maternal and child healthcare" in p1.description.lower()
    assert "React" in p1.technologies
    assert "MongoDB" in p1.technologies

    p2 = c.projects[1]
    assert "CampVerse" in p2.title
    assert "4-member team" in p2.description.lower()
    assert "HTML" in p2.technologies
    assert "JavaScript" in p2.technologies

    # Ensure bogus fragments are not projects
    proj_titles = [p.title.lower() for p in c.projects]
    assert not any("built a full" in t for t in proj_titles)
    assert not any("tech stack" in t for t in proj_titles)
    assert not any("led feature" in t for t in proj_titles)

    # 2. Awards & Honors validation: Real achievements, no skills leakage
    assert len(c.awards) == 2
    assert any("Top 50 Finalist" in a for a in c.awards)
    assert any("Core Member" in a for a in c.awards)
    # Ensure skill headers didn't leak into awards
    assert not any("core cs" in a.lower() for a in c.awards)
    assert not any("databases" in a.lower() for a in c.awards)
    assert not any("tools" in a.lower() for a in c.awards)

    # 3. Skills validation: Real skills, no category headers
    assert "Java" in c.skills
    assert "Python" in c.skills
    assert "Node.js" in c.skills
    assert "PostgreSQL" in c.skills
    assert "MongoDB" in c.skills
    assert "Git" in c.skills
    assert not any(s.lower() in ("languages", "backend", "ai & automation", "core cs", "databases", "tools") for s in c.skills)

    # 4. Certifications validation
    assert any("Oracle Cloud Infrastructure" in cert for cert in c.certifications)
    assert any("Bits and Bytes" in cert for cert in c.certifications)

